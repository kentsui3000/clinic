"""
產出「徐嘉賢診所員工分紅自動化系統」.xlsx 模板
依照 README.md / sheet_structure.md 規格建立 6 張分頁。

執行：
    python3 build_template.py

輸出：
    clinic_bonus_template.xlsx
"""

from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment


OUTPUT_FILE = "clinic_bonus_template.xlsx"


HEADER_FONT = Font(name="Microsoft JhengHei", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5597")
NOTE_FONT = Font(name="Microsoft JhengHei", size=10, italic=True, color="666666")
TITLE_FONT = Font(name="Microsoft JhengHei", size=14, bold=True, color="2F5597")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header_row(ws, row_idx, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_note(ws, cell_ref, text):
    ws[cell_ref].comment = Comment(text, "系統")


# ============================================================
# 分頁 1：員工總表
# ============================================================
def build_employee_sheet(wb):
    ws = wb.create_sheet("員工總表")
    headers = ["員工姓名", "職位", "在職", "加入日期", "離職日期", "備註"]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    initial = [
        ["芯儀", "護理師", "✓", None, None, None],
        ["郁芹", "護理師", "✓", None, None, None],
        ["佩芷", "護理師", "✓", None, None, None],
        ["嘉惠", "行政",   "✓", None, None, None],
        ["春言", "行政",   "✓", None, None, None],
    ]
    for row in initial:
        ws.append(row)

    for r in range(2, 2 + len(initial)):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = CENTER if c in (2, 3) else LEFT
            cell.border = BORDER

    for r in range(2, 2 + len(initial)):
        ws.cell(row=r, column=4).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"

    set_col_widths(ws, [12, 10, 8, 14, 14, 30])
    ws.freeze_panes = "A2"

    add_note(ws, "C1",
        "在職欄填「✓」表示在職。離職時請改為空白，"
        "不要刪除整列（保留歷史分紅明細對應）。"
        "存檔後 Apps Script 會自動同步到兩張 Form 的下拉選項。")

    # 說明列
    ws.cell(row=8, column=1).value = "※ 維護說明"
    ws.cell(row=8, column=1).font = Font(bold=True, color="C00000")
    ws.cell(row=9, column=1).value = "1) 新增員工：新增一列、填姓名、職位，於 C 欄填「✓」。"
    ws.cell(row=10, column=1).value = "2) 員工離職：只把 C 欄改為空白，不要刪除整列。"
    ws.cell(row=11, column=1).value = "3) 此分頁為唯一手動維護處，存檔後自動觸發同步。"
    for r in (9, 10, 11):
        ws.cell(row=r, column=1).font = NOTE_FONT


# ============================================================
# 分頁 2：單價參數
# ============================================================
def build_pricing_sheet(wb):
    ws = wb.create_sheet("單價參數")
    headers = ["項目", "數值", "說明"]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    rows = [
        ["看診人數門檻",     60,  "超過才開始算"],
        ["超門檻單價",        5,  "每人 5 元"],
        ["自費流感單價",      5,  "每支 5 元"],
        ["代謝症候群-收案", 100,  "每筆 100 元"],
        ["代謝症候群-追蹤",  20,  "每筆 20 元"],
        ["值日生津貼",     100,  "每時段 100 元"],
        ["值日生人數門檻",   0,  "（目前未使用，預留）"],
    ]
    for row in rows:
        ws.append(row)

    for r in range(2, 2 + len(rows)):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = LEFT if c == 3 else CENTER
            cell.border = BORDER

    set_col_widths(ws, [22, 10, 36])
    ws.freeze_panes = "A2"

    ws.cell(row=11, column=1).value = "※ 注意"
    ws.cell(row=11, column=1).font = Font(bold=True, color="C00000")
    ws.cell(row=12, column=1).value = (
        "Apps Script V1 目前是 hardcode 上述數字。"
        "若未來想改成從此表讀取，需修改 apps_script.gs。"
    )
    ws.cell(row=12, column=1).font = NOTE_FONT


# ============================================================
# 分頁 3：Form回應_當班
# ============================================================
def build_form_daily_sheet(wb):
    ws = wb.create_sheet("Form回應_當班")
    headers = [
        "時間戳記", "日期", "門診時段", "本時段值日生", "本班當值人員",
        "填寫人", "核對人", "當時段有效看診人數", "自費流感疫苗支數", "備註",
    ]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    set_col_widths(ws, [18, 12, 10, 14, 30, 12, 12, 18, 16, 24])
    ws.freeze_panes = "A2"

    add_note(ws, "A1",
        "此分頁由 Google Form 1（每日當班結帳）自動寫入，請勿手動編輯。"
        "上傳到 Google Sheets 後，將 Form 1 的回應目的地設為這張分頁。")
    add_note(ws, "E1",
        "Google Form 寫入時用「, 」（逗號+空格）分隔複選值。"
        "Apps Script 內 split(', ') 會自動拆人。")


# ============================================================
# 分頁 4：Form回應_代謝
# ============================================================
def build_form_meta_sheet(wb):
    ws = wb.create_sheet("Form回應_代謝")
    headers = ["時間戳記", "日期", "時段", "活動類型", "執行人員", "病歷號", "備註"]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    set_col_widths(ws, [18, 12, 10, 14, 14, 16, 24])
    ws.freeze_panes = "A2"

    add_note(ws, "A1",
        "此分頁由 Google Form 2（代謝症候群活動記錄）自動寫入，請勿手動編輯。")


# ============================================================
# 分頁 5：分紅明細
# ============================================================
def build_detail_sheet(wb):
    ws = wb.create_sheet("分紅明細")
    headers = [
        "寫入時間", "日期", "時段", "員工", "分紅類型",
        "金額", "看診人數", "流感支數", "病歷號",
    ]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    set_col_widths(ws, [18, 12, 10, 12, 14, 10, 12, 12, 14])
    ws.freeze_panes = "A2"

    add_note(ws, "A1",
        "此分頁由 Apps Script 自動 append（每員工每筆獨立一列）。"
        "第 1 列為標題列，第 2 列起由腳本自動寫入；請勿手動編輯。"
        "所有月度結算 QUERY 都從這張表計算。")
    add_note(ws, "G1", "僅「基礎分紅」會填入看診人數；其他類型留空白。")
    add_note(ws, "H1", "僅「基礎分紅」會填入流感支數；其他類型留空白。")


# ============================================================
# 分頁 6：月度結算
# ============================================================
def build_summary_sheet(wb):
    ws = wb.create_sheet("月度結算")

    set_col_widths(ws, [14, 14, 4, 14, 14, 14, 4, 14, 14, 14, 4, 14, 14, 14, 14])

    # 區塊 1：本月各員工總分紅
    ws["A1"] = "區塊 1：本月各員工總分紅"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")
    ws["A2"] = (
        '=QUERY(\'分紅明細\'!B:F,'
        '"SELECT D, SUM(F)'
        " WHERE B >= date '\"&TEXT(EOMONTH(TODAY(),-1)+1,\"yyyy-MM-dd\")&\"'"
        " AND B <= date '\"&TEXT(EOMONTH(TODAY(),0),\"yyyy-MM-dd\")&\"'"
        ' GROUP BY D'
        ' ORDER BY SUM(F) DESC'
        ' LABEL SUM(F) ' + "'本月總分紅'" + '", 1)'
    )

    # 區塊 2：本月各員工分紅明細（含類型拆分）
    ws["D1"] = "區塊 2：本月各員工分紅明細（含類型拆分）"
    ws["D1"].font = TITLE_FONT
    ws.merge_cells("D1:F1")
    ws["D2"] = (
        '=QUERY(\'分紅明細\'!B:F,'
        '"SELECT D, E, SUM(F)'
        " WHERE B >= date '\"&TEXT(EOMONTH(TODAY(),-1)+1,\"yyyy-MM-dd\")&\"'"
        " AND B <= date '\"&TEXT(EOMONTH(TODAY(),0),\"yyyy-MM-dd\")&\"'"
        ' GROUP BY D, E'
        ' ORDER BY D, E'
        ' LABEL SUM(F) ' + "'小計'" + '", 1)'
    )

    # 區塊 3：指定月份查詢（手動）
    ws["H1"] = "查詢年月"
    ws["H1"].font = Font(name="Microsoft JhengHei", size=11, bold=True, color="C00000")
    ws["H2"] = "2026-04"
    ws["H2"].alignment = CENTER
    ws["H2"].border = BORDER
    add_note(ws, "H2", "輸入欲查詢的年月（格式：YYYY-MM）。修改後 I 欄表格會自動更新。")

    ws["I1"] = "區塊 3：指定月份各員工總分紅"
    ws["I1"].font = TITLE_FONT
    ws.merge_cells("I1:J1")
    ws["I2"] = (
        '=QUERY(\'分紅明細\'!B:F,'
        '"SELECT D, SUM(F)'
        " WHERE B >= date '\"&H2&\"-01'"
        " AND B <= date '\"&TEXT(EOMONTH(DATEVALUE(H2&\"-01\"),0),\"yyyy-MM-dd\")&\"'"
        ' GROUP BY D'
        ' ORDER BY SUM(F) DESC'
        ' LABEL SUM(F) ' + "'TARGET月總分紅'" + '", 1)'
    )

    # 區塊 4：本月診次彙總（cross-check 用）
    ws["L1"] = "區塊 4：本月診次彙總（cross-check 帳本）"
    ws["L1"].font = TITLE_FONT
    ws.merge_cells("L1:O1")
    ws["L2"] = (
        '=QUERY(\'Form回應_當班\'!A:J,'
        '"SELECT B, C, H, I'
        " WHERE B >= date '\"&TEXT(EOMONTH(TODAY(),-1)+1,\"yyyy-MM-dd\")&\"'"
        " AND B <= date '\"&TEXT(EOMONTH(TODAY(),0),\"yyyy-MM-dd\")&\"'"
        ' ORDER BY B, C'
        ' LABEL B ' + "'日期', C '時段', H '看診人數', I '流感支數'" + '", 1)'
    )

    ws.freeze_panes = "A3"

    # 底部使用說明
    ws["A20"] = "※ 使用說明"
    ws["A20"].font = Font(bold=True, color="C00000")
    notes = [
        "1) 此 .xlsx 上傳 Google Drive → 右鍵「以 Google 試算表開啟」後，所有 QUERY 公式才會運作。",
        "2) 本機 Excel 不支援 QUERY()，請忽略 #NAME? 錯誤，以 Sheets 為準。",
        "3) 區塊 1、2、4 自動抓「本月」，月初自動切換。",
        "4) 區塊 3 修改 H2 年月即可查詢任意月份。",
        "5) 員工嚴禁存取此檔案；僅 owner（徐醫師）+ 主任（editor）有權限。",
    ]
    for i, text in enumerate(notes, start=21):
        ws.cell(row=i, column=1).value = text
        ws.cell(row=i, column=1).font = NOTE_FONT
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=7)


# ============================================================
# 封面/說明分頁（放最前面）
# ============================================================
def build_cover_sheet(wb):
    ws = wb.active
    ws.title = "說明"

    ws["A1"] = "徐嘉賢診所員工分紅自動化系統"
    ws["A1"].font = Font(name="Microsoft JhengHei", size=18, bold=True, color="2F5597")
    ws.merge_cells("A1:E1")

    ws["A2"] = "版本 v1.0 ／ 規格對齊日 2026-05-13"
    ws["A2"].font = NOTE_FONT
    ws.merge_cells("A2:E2")

    blocks = [
        ("一、上傳到 Google Drive", [
            "1) 把本檔上傳至徐醫師個人 Google Drive。",
            "2) 在 Drive 右鍵此檔 → 開啟方式 → Google 試算表。",
            "3) 重新命名為「徐嘉賢診所分紅系統」。",
        ]),
        ("二、建立兩張 Google Form", [
            "Form 1：每日當班結帳（欄位請完全依 README 規格）。",
            "Form 2：代謝症候群活動記錄。",
            "兩張 Form 的回應目的地皆指向此 Sheets，",
            "Form 1 寫入「Form回應_當班」、Form 2 寫入「Form回應_代謝」。",
        ]),
        ("三、安裝 Apps Script", [
            "擴充功能 → Apps Script，貼上 apps_script.gs 完整內容。",
            "填入兩張 Form 的 ID（編輯網址 /d/ 之後那串）。",
            "設定觸發器：syncEmployeesToForms（編輯時）、",
            "onDailyFormSubmit（Form 1 提交時）、onMetaFormSubmit（Form 2 提交時）。",
            "手動執行一次 syncEmployeesToForms() 完成首次同步。",
        ]),
        ("四、權限設定（極重要）", [
            "整份 Sheets 嚴禁分享給任何員工。",
            "僅徐醫師（owner）+ 主任（可選 editor）有存取權。",
            "員工只接觸 Form 連結，不接觸 Sheets。",
        ]),
        ("五、分紅規則摘要", [
            "基礎分紅 = MAX(0, 看診 - 流感 - 60) × 5 + 流感 × 5（每位當班員工平均分）。",
            "值日生津貼 = 100 元（僅上午/下午、值日生一人）。",
            "代謝症候群：收案 100 元/筆、追蹤 20 元/筆，給該筆執行人員。",
            "詳見 README.md 與 apps_script.gs。",
        ]),
        ("六、分頁清單", [
            "1. 員工總表（唯一手動維護處）",
            "2. 單價參數",
            "3. Form回應_當班（Form 1 自動寫入）",
            "4. Form回應_代謝（Form 2 自動寫入）",
            "5. 分紅明細（Apps Script 自動 append，月結算來源）",
            "6. 月度結算（含 QUERY 公式，上傳 Sheets 後生效）",
        ]),
    ]

    row = 4
    for title, lines in blocks:
        ws.cell(row=row, column=1).value = title
        ws.cell(row=row, column=1).font = Font(
            name="Microsoft JhengHei", size=12, bold=True, color="2F5597"
        )
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        row += 1
        for line in lines:
            c = ws.cell(row=row, column=1)
            c.value = line
            c.font = Font(name="Microsoft JhengHei", size=11)
            c.alignment = LEFT
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            row += 1
        row += 1

    set_col_widths(ws, [22, 22, 22, 22, 22])


# ============================================================
# 主程式
# ============================================================
def main():
    wb = Workbook()
    build_cover_sheet(wb)
    build_employee_sheet(wb)
    build_pricing_sheet(wb)
    build_form_daily_sheet(wb)
    build_form_meta_sheet(wb)
    build_detail_sheet(wb)
    build_summary_sheet(wb)

    wb.save(OUTPUT_FILE)
    print(f"已產出：{OUTPUT_FILE}")
    print(f"分頁：{wb.sheetnames}")


if __name__ == "__main__":
    main()
